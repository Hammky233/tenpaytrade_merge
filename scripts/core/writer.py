"""
Excel 输出模块
"""

import logging
import os
import re
import pandas as pd
from openpyxl.styles import PatternFill
from utils.columns import find_column

logger = logging.getLogger("TenpayMerge")

# openpyxl 非法字符正则（与 openpyxl.cell.cell.ILLEGAL_CHARACTERS_RE 一致）
# 控制字符 \x00-\x08, \x0B-\x0C, \x0E-\x1F 在 Excel/XML 中不允许
_ILLEGAL_CHARS_RE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def _sanitize_for_excel(df: pd.DataFrame) -> pd.DataFrame:
    """
    移除 DataFrame 所有字符串列中的 Excel 非法控制字符。
    兼容 pandas 2.0+ StringDtype 和旧版 object dtype。
    """
    for col in df.columns:
        # pandas 2.0+ 的 dtype=str 是 StringDtype，不是 object，用 is_string_dtype 统一判断
        if pd.api.types.is_string_dtype(df[col]):
            df[col] = df[col].astype(str).str.replace(_ILLEGAL_CHARS_RE, '', regex=True)
    return df

# 黄色填充（用于标记不确定的行）
YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

# 默认隐藏的列（关键词匹配，所有关键词必须同时出现）
HIDDEN_COLUMN_KEYWORDS = [
    ['交易单号'],
    ['大单号'],
    ['借贷', '类型'],
    ['账户余额(元)'],
    ['交易', '业务', '类型'],
    ['用户', '银行卡号'],
    ['用户', '网银联单号'],
    ['网联', '银联'],
    ['第三方账户名称'],
    ['对手方', '银行卡号'],
    ['对手', '银行', '名称'],
    ['对手', '网银联单号'],
    ['基金公司', '信息'],
    ['间联', '非间联'],
    ['对手方', '接收', '时间'],
]

# 固定列宽的列：{关键词元组: 宽度}
FIXED_WIDTH_COLUMNS = {
    ('用户ID',): 7,
    ('对手方ID',): 7,
    ('用户侧账号名称',):8,
    ('交易用途类型',):5,
    ('备注1',):9,
    ('对手方接收金额(元)',):8,
}


def _format_worksheet(
    worksheet,
    df: pd.DataFrame,
    yellow_mask = None,
):
    """
    对工作表应用统一格式：列宽自适应、固定列宽、隐藏列、冻结表头、可选的黄色标记。

    Args:
        worksheet: openpyxl Worksheet 对象
        df: 对应的 DataFrame
        yellow_mask: pd.Series 或 None — True 的行整行标黄（0-based index 对齐 df）
    """
    # 调整列宽（自适应）
    for column in worksheet.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                val = str(cell.value) if cell.value is not None else ''
                # 中文字符按2个字符宽度计算
                char_len = sum(2 if '一' <= c <= '鿿' or '　' <= c <= '〿' else 1 for c in val)
                if char_len > max_length:
                    max_length = char_len
            except Exception:
                pass
        adjusted_width = min(max_length + 2, 60)
        worksheet.column_dimensions[column_letter].width = max(adjusted_width, 8)

    # 应用固定列宽（覆盖自适应结果）
    for col_idx, col_name in enumerate(df.columns, 1):
        column_letter = worksheet.cell(row=1, column=col_idx).column_letter
        col_str = str(col_name)
        for keywords, width in FIXED_WIDTH_COLUMNS.items():
            if all(k in col_str for k in keywords):
                worksheet.column_dimensions[column_letter].width = width
                break

    # 隐藏指定列
    for col_idx, col_name in enumerate(df.columns, 1):
        column_letter = worksheet.cell(row=1, column=col_idx).column_letter
        col_str = str(col_name)
        for keywords in HIDDEN_COLUMN_KEYWORDS:
            if all(k in col_str for k in keywords):
                worksheet.column_dimensions[column_letter].hidden = True
                break

    # 冻结表头行
    worksheet.freeze_panes = 'A2'

    # 黄色标记：根据传入的 mask 整行标黄
    if yellow_mask is not None:
        yellow_indices = set(yellow_mask[yellow_mask].index)
        for row_idx in range(2, worksheet.max_row + 1):
            df_row = row_idx - 2  # Excel 行号 → df index (0-based)
            if df_row in yellow_indices:
                for c in range(1, worksheet.max_column + 1):
                    worksheet.cell(row=row_idx, column=c).fill = YELLOW_FILL

    # 隐藏辅助列（以 _ 开头的内部列）
    for col_idx, col_name in enumerate(df.columns, 1):
        col_str = str(col_name)
        if col_str.startswith('_'):
            column_letter = worksheet.cell(row=1, column=col_idx).column_letter
            worksheet.column_dimensions[column_letter].hidden = True


def write_excel(
    df: pd.DataFrame,
    output_path: str,
    sheet_name: str = "财付通交易汇总",
    parking_df: pd.DataFrame | None = None,
    special_df: pd.DataFrame | None = None,
    mahjong_df: pd.DataFrame | None = None,
    mahjong_stats_df: pd.DataFrame | None = None,
    grp_df: pd.DataFrame | None = None,
    grp_stats_df: pd.DataFrame | None = None,
) -> str:
    """
    将 DataFrame 写入 Excel 文件，自动调整列宽。

    Args:
        df: 待输出的主 DataFrame
        output_path: 输出文件路径（含 .xlsx 扩展名）
        sheet_name: 主工作表名称
        parking_df: 可选的停车缴费 DataFrame，写入独立工作表「停车缴费」
        special_df: 可选的特殊交易 DataFrame，写入独立工作表「特殊交易」
        mahjong_df: 可选的疑似麻友交易明细 DataFrame，写入独立工作表「疑似麻友」
        mahjong_stats_df: 可选的疑似麻友统计 DataFrame，写入独立工作表「疑似麻友-统计」
        grp_df: 可选的群红包交易明细 DataFrame，写入独立工作表「群红包记录」
        grp_stats_df: 可选的群红包统计 DataFrame，写入独立工作表「群红包-统计」

    Returns:
        输出文件路径
    """
    if df.empty:
        logger.warning("DataFrame 为空，不生成输出")
        return ""

    # 确保输出目录存在
    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    # 清除 Excel 非法控制字符（否则 openpyxl 会抛 IllegalCharacterError）
    _sanitize_for_excel(df)
    if parking_df is not None and not parking_df.empty:
        _sanitize_for_excel(parking_df)
    if special_df is not None and not special_df.empty:
        _sanitize_for_excel(special_df)
    if mahjong_df is not None and not mahjong_df.empty:
        _sanitize_for_excel(mahjong_df)
    if mahjong_stats_df is not None and not mahjong_stats_df.empty:
        _sanitize_for_excel(mahjong_stats_df)
    if grp_df is not None and not grp_df.empty:
        _sanitize_for_excel(grp_df)
    if grp_stats_df is not None and not grp_stats_df.empty:
        _sanitize_for_excel(grp_stats_df)

    # 最终返回路径（默认 xlsx）
    result_path = output_path

    try:
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            # 写入主数据表
            df.to_excel(writer, index=False, sheet_name=sheet_name)
            main_ws = writer.sheets[sheet_name]
            _format_worksheet(main_ws, df)

            # ── 以下每个 sheet 独立 try-except，单个失败不影响其他 ──

            # 写入停车缴费工作表
            if parking_df is not None and not parking_df.empty:
                try:
                    parking_sheet_name = "停车缴费"
                    parking_df.to_excel(writer, index=False, sheet_name=parking_sheet_name)
                    parking_ws = writer.sheets[parking_sheet_name]
                    from core.parking import build_parking_yellow_mask
                    yellow_mask = build_parking_yellow_mask(parking_df)
                    _format_worksheet(parking_ws, parking_df, yellow_mask=yellow_mask)
                    logger.info(f"停车缴费工作表已输出: {len(parking_df)} 条记录")
                except Exception as _e:
                    logger.error(f"停车缴费工作表写入失败: {_e}", exc_info=True)

            # 写入特殊交易工作表
            if special_df is not None and not special_df.empty:
                try:
                    special_sheet_name = "特殊交易"
                    special_df.to_excel(writer, index=False, sheet_name=special_sheet_name)
                    special_ws = writer.sheets[special_sheet_name]
                    _format_worksheet(special_ws, special_df)
                    logger.info(f"特殊交易工作表已输出: {len(special_df)} 条记录")
                except Exception as _e:
                    logger.error(f"特殊交易工作表写入失败: {_e}", exc_info=True)

            # 写入疑似麻友工作表（交易明细）
            if mahjong_df is not None and not mahjong_df.empty:
                try:
                    mahjong_sheet_name = "疑似麻友"
                    mahjong_df.to_excel(writer, index=False, sheet_name=mahjong_sheet_name)
                    mahjong_ws = writer.sheets[mahjong_sheet_name]
                    _format_worksheet(mahjong_ws, mahjong_df)
                    logger.info(f"疑似麻友明细工作表已输出: {len(mahjong_df)} 条记录")
                except Exception as _e:
                    logger.error(f"疑似麻友明细工作表写入失败: {_e}", exc_info=True)

            # 写入疑似麻友统计表（独立 sheet）
            if mahjong_stats_df is not None and not mahjong_stats_df.empty:
                try:
                    stats_sheet_name = "疑似麻友-统计"
                    mahjong_stats_df.to_excel(writer, index=False, sheet_name=stats_sheet_name)
                    stats_ws = writer.sheets[stats_sheet_name]
                    _format_worksheet(stats_ws, mahjong_stats_df)
                    logger.info(f"疑似麻友统计工作表已输出: {len(mahjong_stats_df)} 个嫌疑人")
                except Exception as _e:
                    logger.error(f"疑似麻友统计工作表写入失败: {_e}", exc_info=True)

            # 写入群红包记录工作表
            if grp_df is not None and not grp_df.empty:
                try:
                    grp_sheet_name = "群红包记录"
                    grp_df.to_excel(writer, index=False, sheet_name=grp_sheet_name)
                    grp_ws = writer.sheets[grp_sheet_name]
                    _format_worksheet(grp_ws, grp_df)
                    logger.info(f"群红包记录工作表已输出: {len(grp_df)} 条记录")
                except Exception as _e:
                    logger.error(f"群红包记录工作表写入失败: {_e}", exc_info=True)

            # 写入群红包统计工作表
            if grp_stats_df is not None and not grp_stats_df.empty:
                try:
                    grp_stats_sheet_name = "群红包-统计"
                    grp_stats_df.to_excel(writer, index=False, sheet_name=grp_stats_sheet_name)
                    grp_stats_ws = writer.sheets[grp_stats_sheet_name]
                    _format_worksheet(grp_stats_ws, grp_stats_df)
                    logger.info(f"群红包统计工作表已输出: {len(grp_stats_df)} 个对手方")
                except Exception as _e:
                    logger.error(f"群红包统计工作表写入失败: {_e}", exc_info=True)

        logger.info(f"Excel 输出完成: {output_path}")
        return result_path

    except Exception as e:
        # 仅当「主 sheet」写入失败时才走到这里（子 sheet 错误已在上面独立处理）
        import traceback
        err_detail = traceback.format_exc()
        logger.error(f"Excel 输出失败(主表): {e}\n{err_detail}")
        # fallback: 输出 CSV
        csv_path = output_path.replace('.xlsx', '.csv')
        df.to_csv(csv_path, index=False, encoding='utf-8-sig')
        logger.info(f"已输出 CSV 备选: {csv_path}")
        return csv_path


# ── 注册信息专用隐藏列 ──────────────────────────────────────────────

REG_HIDDEN_COLUMN_KEYWORDS = [
    ['开户行信息'],
    ['银行账号'],
    ['数据来源'],
    ['调证编号'],
]

REG_FIXED_WIDTH_COLUMNS = {
    ('账户状态',): 8,
    ('账号',): 18,
    ('注册身份证号',): 20,
    ('变更类型',): 12,
    ('当前状态',): 8,
}


def write_reg_excel(
    basic_df: pd.DataFrame,
    changes_df: pd.DataFrame,
    output_path: str,
    person_df: pd.DataFrame | None = None,
) -> str:
    """
    将注册信息写入 Excel 文件（三 sheet：「注册信息汇总」+「变更记录」+「基础信息」）。

    Args:
        basic_df: 注册信息汇总 DataFrame
        changes_df: 变更记录 DataFrame
        output_path: 输出文件路径（含 .xlsx 扩展名）
        person_df: 可选的人员基础信息 DataFrame（注册姓名 + 身份证号 + 手机 去重）

    Returns:
        输出文件路径
    """
    if basic_df.empty and changes_df.empty and (person_df is None or person_df.empty):
        logger.warning("注册信息 DataFrame 为空，不生成输出")
        return ""

    # 确保输出目录存在
    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    try:
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            # ── Sheet 1: 注册信息汇总 ──
            if not basic_df.empty:
                basic_df.to_excel(writer, index=False, sheet_name="注册信息汇总")
                basic_ws = writer.sheets["注册信息汇总"]
                _format_worksheet(basic_ws, basic_df)
                _apply_reg_fixed_widths(basic_ws, basic_df)
                _apply_reg_hidden_columns(basic_ws, basic_df)

            # ── Sheet 2: 变更记录 ──
            if not changes_df.empty:
                changes_df.to_excel(writer, index=False, sheet_name="变更记录")
                changes_ws = writer.sheets["变更记录"]
                _format_worksheet(changes_ws, changes_df)
                _apply_reg_fixed_widths(changes_ws, changes_df)
                _apply_reg_hidden_columns(changes_ws, changes_df)

            # ── Sheet 3: 基础信息（自然人级别去重）──
            if person_df is not None and not person_df.empty:
                person_df.to_excel(writer, index=False, sheet_name="基础信息")
                person_ws = writer.sheets["基础信息"]
                _format_worksheet(person_ws, person_df)
                _apply_reg_fixed_widths(person_ws, person_df)

        logger.info(f"注册信息 Excel 输出完成: {output_path}")
        return output_path

    except Exception as e:
        logger.error(f"注册信息 Excel 输出失败: {e}")
        # fallback: 输出 CSV
        if not basic_df.empty:
            csv_path = output_path.replace('.xlsx', '_汇总.csv')
            basic_df.to_csv(csv_path, index=False, encoding='utf-8-sig')
            logger.info(f"已输出 CSV 备选: {csv_path}")
        if not changes_df.empty:
            csv_path = output_path.replace('.xlsx', '_变更.csv')
            changes_df.to_csv(csv_path, index=False, encoding='utf-8-sig')
            logger.info(f"已输出 CSV 备选: {csv_path}")
        return ""


def _apply_reg_hidden_columns(worksheet, df: pd.DataFrame):
    """隐藏注册信息专用敏感/冗余列。"""
    for col_idx, col_name in enumerate(df.columns, 1):
        column_letter = worksheet.cell(row=1, column=col_idx).column_letter
        col_str = str(col_name)
        for keywords in REG_HIDDEN_COLUMN_KEYWORDS:
            if all(k in col_str for k in keywords):
                worksheet.column_dimensions[column_letter].hidden = True
                break


def _apply_reg_fixed_widths(worksheet, df: pd.DataFrame):
    """覆盖注册信息专用固定列宽。"""
    for col_idx, col_name in enumerate(df.columns, 1):
        column_letter = worksheet.cell(row=1, column=col_idx).column_letter
        col_str = str(col_name)
        for keywords, width in REG_FIXED_WIDTH_COLUMNS.items():
            if all(k in col_str for k in keywords):
                worksheet.column_dimensions[column_letter].width = width
                break
