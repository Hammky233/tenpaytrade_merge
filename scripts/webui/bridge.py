"""
pywebview Bridge API — Python 端暴露给前端的方法
"""

import os
import sys
import threading

# tkinter 在部分 Linux 发行版需要单独安装 python3-tk
try:
    from tkinter import filedialog
    _TKINTER_AVAILABLE = True
except ImportError:
    _TKINTER_AVAILABLE = False

# 确保项目根目录在 sys.path 中
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# 初始化日志（GUI 模式下必须手动配置，否则 writer 等模块的 logger 输出不可见）
from utils.logger import setup_logger as _setup_logger
_setup_logger(log_dir=None)  # 控制台 handler，GUI 模式下不写文件

from service.pipeline import TenpayPipeline, ProgressInfo, post_merge_analysis
from core.merger import merge_dataframes, deduplicate
from core.writer import write_excel
from core.reg_reader import read_tenpay_reg_info
from core.reg_processor import process_reg_data, build_person_info
from core.writer import write_reg_excel
import pandas as pd


class Api:
    """pywebview JS-Python Bridge"""

    def __init__(self):
        self._pipeline: TenpayPipeline | None = None
        self._thread: threading.Thread | None = None
        self._api_key: str = ""  # 仅内存存储，GUI 关闭即清空

    def get_version(self) -> str:
        """返回当前版本号"""
        from version import VERSION
        return VERSION

    def get_author(self) -> str:
        """返回工具制作人"""
        from version import AUTHOR
        return AUTHOR

    def select_folder(self) -> str:
        """打开文件夹选择对话框，返回所选路径（取消返回空字符串）"""
        if not _TKINTER_AVAILABLE:
            return self._fallback_dialog_error()
        import tkinter as tk
        root = tk.Tk()
        root.withdraw()
        root.attributes('-topmost', True)
        folder = filedialog.askdirectory(title="选择文件夹")
        root.destroy()
        return folder or ""

    def select_file(self, file_types: str = "Excel files (*.xlsx)|*.xlsx") -> str:
        """打开文件选择对话框，返回所选路径（取消返回空字符串）"""
        if not _TKINTER_AVAILABLE:
            return self._fallback_dialog_error()
        import tkinter as tk
        root = tk.Tk()
        root.withdraw()
        root.attributes('-topmost', True)
        filepath = filedialog.askopenfilename(
            title="选择文件",
            filetypes=[(ft.split('|')[0], ft.split('|')[1]) for ft in file_types.split(';')] if file_types else None,
        )
        root.destroy()
        return filepath or ""

    def select_files(self, file_types: str = "Excel files (*.xlsx)|*.xlsx") -> str:
        """多文件选择，返回用 | 分隔的路径字符串"""
        if not _TKINTER_AVAILABLE:
            return self._fallback_dialog_error()
        import tkinter as tk
        root = tk.Tk()
        root.withdraw()
        root.attributes('-topmost', True)
        files = filedialog.askopenfilenames(
            title="选择文件（可多选）",
            filetypes=[(ft.split('|')[0], ft.split('|')[1]) for ft in file_types.split(';')] if file_types else None,
        )
        root.destroy()
        return "|".join(files) if files else ""

    def _fallback_dialog_error(self) -> str:
        """tkinter 不可用时的错误提示（Linux 需安装 python3-tk）"""
        import platform
        msg = "文件对话框不可用：tkinter 未安装"
        if platform.system() == "Linux":
            msg += "。请运行: sudo apt install python3-tk (或等效命令)"
        return f"__ERROR__:{msg}"

    def start_batch_process(self, source: str, output: str, output_name: str = "Tenpay_merge.xlsx",
                            process_reg: bool = False, timestamp: str = "",
                            enable_location: bool = False, api_key: str = "") -> str:
        """
        启动单批次处理（后台线程），返回 "started" 或错误信息

        Args:
            source: 数据源文件夹
            output: 输出文件夹
            output_name: 交易流水输出文件名
            process_reg: 是否同时清洗注册信息（TenpayRegInfo.txt）
            timestamp: 时间戳（MMDD_HHmm 格式），用于文件名后缀
            enable_location: 是否启用 AI 地点识别
            api_key: DeepSeek API Key（仅内存保存，不持久化）
        """
        if self._thread and self._thread.is_alive():
            return "已有任务正在运行"
        if not source:
            return "请先选择数据源文件夹"
        if not output:
            return "请先选择输出文件夹"

        # 时间戳已由前端拼入 output_name，此处不再重复追加
        # （timestamp 参数仅用于注册信息输出文件名）

        self._pipeline = TenpayPipeline(
            source_dir=source,
            output_dir=output,
            output_name=output_name,
            api_key=api_key if enable_location else None,
        )

        def _run():
            # ── 交易流水清洗 ──
            try:
                self._pipeline.run()
            except Exception as e:
                import traceback
                self._pipeline.progress.status = "error"
                self._pipeline.progress.add_log(f"❌ 处理异常: {e}")
                for line in traceback.format_exc().splitlines():
                    if line.strip():
                        self._pipeline.progress.add_log(f"   {line.strip()}")

            # ── 注册信息清洗（可选）──
            if process_reg and self._pipeline.progress.status != "error":
                try:
                    # 前端在 status="done" 时停止轮询，必须保持 running 状态
                    self._pipeline.progress.status = "running"
                    self._run_reg_process(source, output, timestamp)
                except Exception as e:
                    import traceback
                    self._pipeline.progress.add_log(f"❌ 注册信息清洗异常: {e}")
                    for line in traceback.format_exc().splitlines():
                        if line.strip():
                            self._pipeline.progress.add_log(f"   {line.strip()}")
                finally:
                    # 全部完成后标记 done，前端停止轮询
                    self._pipeline.progress.status = "done"

        self._thread = threading.Thread(target=_run, daemon=True)
        self._thread.start()
        return "started"

    def _run_reg_process(self, source: str, output: str, timestamp: str = ""):
        """在交易流水处理完成后，额外执行注册信息清洗（复用 progress 日志）。"""
        import time
        progress = self._pipeline.progress if self._pipeline else None

        # 注册信息输出文件名
        reg_name = f"TenpayRegInfo_merge_{timestamp}.xlsx" if timestamp else "TenpayRegInfo_merge.xlsx"

        def _log(msg: str):
            if progress:
                progress.add_log(msg)
            # 也走标准日志
            import logging
            logging.getLogger("TenpayMerge").info(msg)

        _log("─" * 40)
        _log("📋 开始清洗注册信息...")

        # 1. 扫描 TenpayRegInfo.txt 文件
        from utils.paths import find_files_by_name
        reg_files = find_files_by_name(source, "TenpayRegInfo.txt")

        if not reg_files:
            _log("⚠️ 未找到 TenpayRegInfo.txt 文件，跳过注册信息清洗")
            return

        _log(f"扫描到 {len(reg_files)} 个 TenpayRegInfo.txt 文件")

        t_start = time.time()

        # 2. 逐文件解析
        success = 0
        fail = 0
        skipped = 0
        records = []

        for i, filepath in enumerate(reg_files, 1):
            try:
                rel_path = os.path.relpath(filepath, source)
            except ValueError:
                rel_path = filepath

            try:
                result = read_tenpay_reg_info(filepath)
                if result is None:
                    skipped += 1
                else:
                    records.append(result)
                    success += 1
            except Exception as e:
                fail += 1
                _log(f"[{i}/{len(reg_files)}] 失败: {rel_path} — {e}")

        # 3. 清洗合并
        basic_df, changes_df = process_reg_data(records)

        # 4. 提取人员基础信息
        person_df = build_person_info(basic_df)

        # 5. 输出 Excel
        output_path = os.path.join(output, reg_name)
        result_path = write_reg_excel(basic_df, changes_df, output_path, person_df=person_df)

        elapsed = time.time() - t_start

        # 6. 日志摘要 + 写入 result 供前端展示
        _log(f"✅ 注册信息清洗完成!")
        _log(f"文件总数: {len(reg_files)} | 成功: {success} | 失败: {fail} | 跳过: {skipped}")
        _log(f"注册信息汇总: {len(basic_df)} 条 | 变更记录: {len(changes_df)} 条 | 基础信息(自然人): {len(person_df)} 人")
        _log(f"耗时: {elapsed:.1f} 秒")
        if result_path:
            _log(f"输出文件: {result_path}")
        else:
            _log("⚠️ 注册信息输出为空（无有效数据）")

        # 写入 reg 结果供前端分别展示
        if progress:
            progress.result["reg_output"] = result_path or ""
            progress.result["reg_basic_rows"] = len(basic_df)
            progress.result["reg_changes_rows"] = len(changes_df)
            progress.result["reg_person_rows"] = len(person_df)
            progress.result["reg_files"] = len(reg_files)
            progress.result["reg_success"] = success
            progress.result["reg_fail"] = fail
            progress.result["reg_elapsed"] = round(elapsed, 1)

    def get_batch_status(self) -> dict:
        """获取单批次处理进度"""
        if self._pipeline:
            return self._pipeline.progress.to_dict()
        return {"status": "idle", "logs": []}

    def start_merge_process(self, input_files_str: str, output: str,
                            enable_location: bool = False, api_key: str = "") -> str:
        """
        启动多批次合并（后台线程）
        input_files_str: 用 | 分隔的文件路径
        enable_location: 是否启用 AI 地点识别
        api_key: DeepSeek API Key
        """
        if self._thread and self._thread.is_alive():
            return "已有任务正在运行"
        if not input_files_str:
            return "请先选择待合并的文件"

        files = [f for f in input_files_str.split("|") if f.strip()]
        if not files:
            return "文件列表为空"

        # 使用 pipeline 的 progress 来传递状态
        self._pipeline = TenpayPipeline(
            source_dir="",  # 合并模式不需要 source
            output_dir=os.path.dirname(output) or ".",
            output_name=os.path.basename(output),
        )
        self._pipeline.progress.total = len(files)

        def _run():
            progress = self._pipeline.progress
            progress.status = "running"
            progress.add_log(f"开始合并 {len(files)} 个批次文件...")

            try:
                dfs = []
                for i, f in enumerate(files):
                    progress.current = i + 1
                    fname = os.path.basename(f)
                    progress.add_log(f"读取: {fname}")
                    try:
                        df = pd.read_excel(f, sheet_name="财付通交易汇总", dtype=str)
                    except ValueError as e:
                        if "not found" in str(e) or "Worksheet named" in str(e):
                            progress.add_log(f"❌ 文件缺少必要工作表「财付通交易汇总」: {fname}")
                        else:
                            progress.add_log(f"❌ 读取失败: {fname} — {e}")
                        progress.fail += 1
                        continue
                    dfs.append(df)
                    progress.add_log(f"  → {len(df)} 行")
                    progress.success += 1

                if not dfs:
                    progress.status = "error"
                    progress.add_log("❌ 所有输入文件均缺少必要工作表「财付通交易汇总」，无法合并")
                    return

                progress.add_log("合并中...")
                merged = merge_dataframes(dfs)
                before = len(merged)

                merged = deduplicate(merged)
                after = len(merged)

                # ── 读取既往「停车缴费」sheet 中的地点映射 ──
                existing_location_map = {}
                if enable_location and api_key:
                    from utils.columns import find_column as _find_col
                    for f in files:
                        try:
                            xl = pd.ExcelFile(f)
                            if "停车缴费" in xl.sheet_names:
                                pf = pd.read_excel(f, sheet_name="停车缴费", dtype=str)
                                # 确保有所需列
                                col_n2 = _find_col(pf.columns, ["备注2"])
                                col_loc = _find_col(pf.columns, ["地点"])
                                if col_n2 and col_loc:
                                    for _, row in pf.iterrows():
                                        loc_val = str(row[col_loc]) if pd.notna(row[col_loc]) else ""
                                        note_val = str(row[col_n2]) if pd.notna(row[col_n2]) else ""
                                        if loc_val and loc_val != "无" and loc_val != "nan" and note_val:
                                            # 保留第一个出现的地点映射
                                            if note_val not in existing_location_map:
                                                existing_location_map[note_val] = loc_val
                            xl.close()
                        except Exception:
                            pass  # 单个文件读取失败不影响整体
                    if existing_location_map:
                        progress.add_log(f"从既往文件读取 {len(existing_location_map)} 条地点映射")

                # 停车缴费识别 + 特殊交易筛选 + 疑似麻友识别 + 群红包识别（通过共享函数，与 pipeline 行为一致）
                analysis = post_merge_analysis(
                    merged,
                    api_key=api_key if enable_location else None,
                    existing_location_map=existing_location_map if existing_location_map else None,
                )
                parking_df = analysis["parking_df"]
                special_df = analysis["special_df"]
                mahjong_df = analysis.get("mahjong_df")
                mahjong_stats_df = analysis.get("mahjong_stats_df")
                grp_df = analysis.get("grp_df")
                grp_stats_df = analysis.get("grp_stats_df")

                write_excel(merged, output,
                            parking_df=parking_df, special_df=special_df,
                            mahjong_df=mahjong_df, mahjong_stats_df=mahjong_stats_df,
                            grp_df=grp_df, grp_stats_df=grp_stats_df)

                progress.status = "done"
                progress.add_log(f"✅ 合并完成!")
                progress.add_log(f"合并前: {before} 行 → 去重后: {after} 行 (移除 {before - after} 条)")
                progress.result = {"output": output, "rows": after, "removed": before - after}

            except Exception as e:
                import traceback
                progress.status = "error"
                progress.add_log(f"❌ 合并失败: {e}")
                for line in traceback.format_exc().splitlines():
                    if line.strip():
                        progress.add_log(f"   {line.strip()}")

        self._thread = threading.Thread(target=_run, daemon=True)
        self._thread.start()
        return "started"

    def get_file_info(self, filepath: str) -> dict:
        """获取单个 xlsx 文件的基本信息（行数、列数）"""
        try:
            df = pd.read_excel(filepath, nrows=0)
            cols = len(df.columns)
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True)
            ws = wb.active
            rows = ws.max_row - 1  # 减去表头
            wb.close()
            return {"rows": rows, "cols": cols, "filename": os.path.basename(filepath)}
        except Exception as e:
            return {"rows": 0, "cols": 0, "filename": os.path.basename(filepath), "error": str(e)}

    def stop_process(self) -> str:
        """停止当前任务（仅标记，线程会自然结束）"""
        if self._pipeline:
            self._pipeline.progress.status = "idle"
            return "stopped"
        return "no task running"

    def get_parking_config(self) -> dict:
        """获取停车缴费识别配置"""
        from core.parking import load_parking_config
        try:
            return load_parking_config()
        except Exception as e:
            return {"error": str(e)}

    def save_parking_config(self, config: dict) -> str:
        """保存停车缴费识别配置到 JSON 文件，返回 "ok" 或错误信息"""
        import json

        # 验证必填字段
        required_fields = ["备注2关键词", "排除关键词", "对手侧账户名称关键词", "车牌省份简称"]
        for field in required_fields:
            if field not in config:
                return f"缺少必填字段: {field}"
            if not isinstance(config[field], list):
                return f"字段 {field} 必须是数组"

        try:
            from utils.paths import get_user_config_dir
            config_dir = get_user_config_dir()
            os.makedirs(config_dir, exist_ok=True)
            config_path = os.path.join(config_dir, "parking_config.json")
            with open(config_path, "w", encoding="utf-8") as f:
                json.dump(config, f, ensure_ascii=False, indent=2)
            return "ok"
        except Exception as e:
            return f"保存失败: {e}"

    def get_time_period_config(self) -> dict:
        """获取时段分类配置"""
        from core.processor import load_time_period_config
        try:
            return load_time_period_config()
        except Exception as e:
            return {"error": str(e)}

    def save_time_period_config(self, config: dict) -> str:
        """保存时段分类配置到 JSON 文件，返回 "ok" 或错误信息"""
        import json

        # 验证必填字段
        if "时段" not in config or not isinstance(config["时段"], list):
            return "缺少必填字段: 时段"
        for period in config["时段"]:
            if not all(k in period for k in ("name", "start", "end")):
                return "每个时段必须包含 name, start, end"

        try:
            from utils.paths import get_user_config_dir
            config_dir = get_user_config_dir()
            os.makedirs(config_dir, exist_ok=True)
            config_path = os.path.join(config_dir, "time_period_config.json")
            with open(config_path, "w", encoding="utf-8") as f:
                json.dump(config, f, ensure_ascii=False, indent=2)
            return "ok"
        except Exception as e:
            return f"保存失败: {e}"

    def re_extract_locations(self, excel_path: str, api_key: str) -> str:
        """
        对已有 Excel 文件的「停车缴费」sheet 独立重新提取地点。

        - 只更新「停车缴费」sheet 的「地点」列，其他 sheet 原样保留
        - 已有非"无"的地点值不覆盖（保护人工修正）

        Returns: 成功时返回 JSON {"status": "ok", "total": N, "before": N, "after": N, "newly": N}
                 失败时返回 JSON {"status": "error", "message": "..."}
        """
        import json

        if not excel_path or not os.path.isfile(excel_path):
            return json.dumps({"status": "error", "message": "文件不存在"}, ensure_ascii=False)
        if not api_key:
            return json.dumps({"status": "error", "message": "请先填入 API Key"}, ensure_ascii=False)

        try:
            from core.location import extract_locations
            from utils.columns import find_column

            # ── 1. 读取所有 sheet ──
            all_sheets = pd.read_excel(excel_path, sheet_name=None, dtype=str)

            if "停车缴费" not in all_sheets:
                return json.dumps({
                    "status": "error",
                    "message": "该文件中未找到「停车缴费」工作表"
                }, ensure_ascii=False)

            parking_df = all_sheets["停车缴费"]

            # ── 2. 确认「备注2」列存在 ──
            col_note2 = find_column(parking_df.columns, ["备注2"])
            if not col_note2:
                return json.dumps({
                    "status": "error",
                    "message": "「停车缴费」sheet 中未找到「备注2」列"
                }, ensure_ascii=False)

            # ── 3. 统计前值 ──
            before_count = 0
            if "地点" in parking_df.columns:
                before_count = int((
                    parking_df["地点"].notna() &
                    (parking_df["地点"] != "无") &
                    (parking_df["地点"] != "")
                ).sum())

            # ── 4. 执行提取 ──
            logs: list[str] = []

            def _log_progress(msg: str):
                import logging
                logging.getLogger("TenpayMerge").info(msg)
                logs.append(msg)

            parking_df = extract_locations(
                parking_df, api_key,
                progress_callback=_log_progress,
            )

            after_count = int((
                parking_df["地点"].notna() &
                (parking_df["地点"] != "无") &
                (parking_df["地点"] != "")
            ).sum())

            newly = after_count - before_count

            # ── 5. 写回 Excel（原地更新，仅操作「停车缴费」sheet 的「地点」列）──
            from core.writer import _sanitize_for_excel, YELLOW_FILL
            from core.parking import build_parking_yellow_mask
            import openpyxl

            _sanitize_for_excel(parking_df)
            wb = openpyxl.load_workbook(excel_path)
            ws = wb["停车缴费"]

            # 5a. 定位「地点」列索引 — 按表头名称查找
            location_col_idx = None
            for cell in ws[1]:
                if cell.value == "地点":
                    location_col_idx = cell.column
                    break

            location_col_found = location_col_idx is not None

            # 5b. 若「地点」列不存在，在表最右侧新建
            if not location_col_found:
                location_col_idx = (ws.max_column or 0) + 1
                ws.cell(row=1, column=location_col_idx).value = "地点"

            # 5c. 逐行只更新「地点」单元格的值（不修改 fill，保留原有格式）
            for df_idx in range(len(parking_df)):
                excel_row = df_idx + 2  # 第 1 行是表头
                loc_value = parking_df.iloc[df_idx]["地点"]
                cell = ws.cell(row=excel_row, column=location_col_idx)
                cell.value = None if pd.isna(loc_value) else loc_value

            # 5d. 仅当「地点」列为新建时，才将原黄色标记扩展到新列对应的单元格
            #     已有列的原黄色标记已在初始创建时由 _format_worksheet 设置，不动
            if not location_col_found:
                yellow_mask = build_parking_yellow_mask(parking_df)
                if yellow_mask is not None:
                    yellow_indices = set(yellow_mask[yellow_mask].index)
                    for df_idx in yellow_indices:
                        excel_row = df_idx + 2
                        ws.cell(row=excel_row, column=location_col_idx).fill = YELLOW_FILL

            wb.save(excel_path)
            wb.close()

            logs.append(
                f"重新提取完成: {after_count}/{len(parking_df)} 条有地点信息"
                f"（新增 {max(0, newly)} 条）"
            )

            return json.dumps({
                "status": "ok",
                "total": len(parking_df),
                "before": before_count,
                "after": after_count,
                "newly": max(0, newly),
                "logs": logs[-10:],
            }, ensure_ascii=False)

        except Exception as e:
            import traceback
            import logging
            logging.getLogger("TenpayMerge").error(
                f"重新提取地点失败: {e}\n{traceback.format_exc()}"
            )
            return json.dumps({
                "status": "error",
                "message": str(e),
            }, ensure_ascii=False)

