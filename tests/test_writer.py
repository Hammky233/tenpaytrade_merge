"""
Excel 输出测试 — scripts/core/writer.py

测试 write_excel() 和 _format_worksheet() 的格式行为。
"""

import os
import tempfile
import pandas as pd
import pytest
from openpyxl import load_workbook
from core.writer import write_excel, _format_worksheet, HIDDEN_COLUMN_KEYWORDS, YELLOW_FILL


class TestWriteExcel:
    """write_excel() 输出行为。"""

    def test_empty_df_returns_empty_string(self):
        """空 DataFrame → 返回空字符串。"""
        df = pd.DataFrame()
        result = write_excel(df, "dummy.xlsx")
        assert result == ""

    def test_basic_output(self):
        """基本输出：主工作表名称为"财付通交易汇总"。"""
        df = pd.DataFrame({"用户ID": ["U001"], "金额": [100.0]})
        with tempfile.TemporaryDirectory() as tmpdir:
            path = os.path.join(tmpdir, "test.xlsx")
            result = write_excel(df, path)
            assert result == path
            assert os.path.exists(path)
            wb = load_workbook(path)
            assert "财付通交易汇总" in wb.sheetnames
            wb.close()

    def test_hidden_columns(self):
        """
        隐藏列：交易单号、大单号等列应被隐藏。
        """
        df = pd.DataFrame({
            "用户ID": ["U001"],
            "交易单号": ["T001"],
            "大单号": ["B001"],
            "金额": [100.0],
        })
        with tempfile.TemporaryDirectory() as tmpdir:
            path = os.path.join(tmpdir, "test.xlsx")
            write_excel(df, path)
            wb = load_workbook(path)
            ws = wb["财付通交易汇总"]

            # 找到各列的位置
            col_map = {}
            for cell in ws[1]:
                col_map[cell.value] = cell.column_letter

            # 交易单号应被隐藏
            assert col_map["交易单号"] is not None
            assert ws.column_dimensions[col_map["交易单号"]].hidden is True

            # 大单号应被隐藏
            assert col_map["大单号"] is not None
            assert ws.column_dimensions[col_map["大单号"]].hidden is True

            # 用户ID 不应被隐藏
            assert col_map["用户ID"] is not None
            hidden = ws.column_dimensions[col_map["用户ID"]].hidden
            assert hidden is None or hidden is False

            wb.close()

    def test_helper_column_hidden(self):
        """以 _ 开头的辅助列应被隐藏。"""
        df = pd.DataFrame({
            "用户ID": ["U001"],
            "_内部标记": ["是"],
            "_备注含省份简称": ["否"],
            "金额": [100.0],
        })
        with tempfile.TemporaryDirectory() as tmpdir:
            path = os.path.join(tmpdir, "test.xlsx")
            write_excel(df, path)
            wb = load_workbook(path)
            ws = wb["财付通交易汇总"]

            col_map = {}
            for cell in ws[1]:
                col_map[cell.value] = cell.column_letter

            # 辅助列应隐藏
            assert ws.column_dimensions[col_map["_内部标记"]].hidden is True
            assert ws.column_dimensions[col_map["_备注含省份简称"]].hidden is True

            # 正常列不应隐藏
            hidden = ws.column_dimensions[col_map["金额"]].hidden
            assert hidden is None or hidden is False

            wb.close()

    def test_freeze_panes(self):
        """表头应冻结（A2）。"""
        df = pd.DataFrame({"用户ID": ["U001"], "金额": [100.0]})
        with tempfile.TemporaryDirectory() as tmpdir:
            path = os.path.join(tmpdir, "test.xlsx")
            write_excel(df, path)
            wb = load_workbook(path)
            ws = wb["财付通交易汇总"]
            assert ws.freeze_panes == "A2"
            wb.close()

    def test_parking_yellow_highlight(self):
        """
        停车缴费标黄：车牌="无"且_备注含省份简称="是"的行被标黄，
        有车牌的行不标黄。
        """
        parking_df = pd.DataFrame({
            "车牌": ["无", "京A12345", "无"],
            "_备注含省份简称": ["是", "否", "否"],
            "金额": [10.0, 20.0, 30.0],
        })

        with tempfile.TemporaryDirectory() as tmpdir:
            path = os.path.join(tmpdir, "test_parking.xlsx")

            # 直接用 write_excel 输出停车场工作表
            from core.parking import build_parking_yellow_mask
            yellow_mask = build_parking_yellow_mask(parking_df)
            write_excel(
                pd.DataFrame({"用户ID": ["U001"]}),
                path,
                parking_df=parking_df,
            )

            wb = load_workbook(path)
            ws = wb["停车缴费"]

            # row 1 = 表头, row 2 = 第1行(无+是 → 标黄)
            # row 2: 车牌="无", _备注含省份简称="是" → 应标黄
            fill_2 = ws.cell(row=2, column=1).fill
            # openpyxl 存颜色为 aarrggbb 格式，黄色 FFFF00 对应 00FFFF00
            assert "FFFF00" in str(fill_2.start_color.rgb).upper()

            # row 3: 车牌="京A12345" → 不标黄
            fill_3 = ws.cell(row=3, column=1).fill
            # 默认填充的 rgb 是 "00000000"，不应包含 FFFF00
            assert "FFFF00" not in str(fill_3.start_color.rgb).upper()

            # row 4: 车牌="无"但省份简称="否" → 不标黄
            fill_4 = ws.cell(row=4, column=1).fill
            assert "FFFF00" not in str(fill_4.start_color.rgb).upper()

            wb.close()

    def test_extra_sheets_not_empty(self):
        """附加工作表（停车缴费）应正确写入数据。"""
        main_df = pd.DataFrame({"用户ID": ["U001"]})
        parking_df = pd.DataFrame({"车牌": ["京A12345"], "金额": [10.0]})

        with tempfile.TemporaryDirectory() as tmpdir:
            path = os.path.join(tmpdir, "test.xlsx")
            write_excel(main_df, path, parking_df=parking_df)

            wb = load_workbook(path)
            assert "停车缴费" in wb.sheetnames
            ws = wb["停车缴费"]
            # 表头+1行数据
            assert ws.max_row == 2
            wb.close()
